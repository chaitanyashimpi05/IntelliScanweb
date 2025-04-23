from flask import Flask, request, jsonify, send_file
from flask_cors import CORS
import os
import pdfplumber
import docx
import spacy
from fpdf import FPDF
from werkzeug.utils import secure_filename
import tempfile

# Initialize Flask app
app = Flask(__name__)
CORS(app)

# Load Spacy NLP model
nlp = spacy.load("en_core_web_sm")

# Function to extract text from PDF
def extract_text_from_pdf(file_path):
    with pdfplumber.open(file_path) as pdf:
        return "\n".join([page.extract_text() for page in pdf.pages if page.extract_text()])

# Function to extract text from DOCX
def extract_text_from_docx(file_path):
    doc = docx.Document(file_path)
    return "\n".join([para.text for para in doc.paragraphs])

# Function to generate PDF report
def generate_pdf_report(job_role, matched_skills, missing_skills):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", "B", 16)
    pdf.cell(200, 10, "IntelliScan Resume Analysis Report", ln=True, align="C")

    pdf.set_font("Arial", "", 12)
    pdf.cell(200, 10, f"Job Role: {job_role}", ln=True)
    pdf.cell(200, 10, f"Matched Skills: {', '.join(matched_skills)}", ln=True)
    pdf.cell(200, 10, f"Missing Skills: {', '.join(missing_skills)}", ln=True)

    temp_file = tempfile.NamedTemporaryFile(delete=False, suffix=".pdf")
    pdf.output(temp_file.name)
    return temp_file.name

# Route to handle file upload and analysis
@app.route('/analyze', methods=['POST'])
def analyze_resume():
    job_role = request.form.get('job_role')
    uploaded_file = request.files.get('resume')

    if not uploaded_file:
        return jsonify({"error": "No file uploaded"})

    filename = secure_filename(uploaded_file.filename)
    if not os.path.exists("uploads"):
        os.makedirs("uploads")
    file_path = os.path.join("uploads", filename)
    uploaded_file.save(file_path)

    if filename.endswith(".pdf"):
        resume_text = extract_text_from_pdf(file_path)
    elif filename.endswith(".docx"):
        resume_text = extract_text_from_docx(file_path)
    else:
        return jsonify({"error": "Unsupported file format"})

    matched_skills = ["Python", "Machine Learning"]
    missing_skills = ["JavaScript", "SQL"]

    pdf_file_path = generate_pdf_report(job_role, matched_skills, missing_skills)
    return send_file(pdf_file_path, as_attachment=True, download_name="resume_analysis_report.pdf")

if __name__ == '__main__':
    app.run(debug=True)
